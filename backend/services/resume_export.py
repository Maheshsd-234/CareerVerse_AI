"""ResumeExporter service: converts resume data into PDF (styled & ATS-optimized), DOCX, and TXT."""

import io
import os
import logging
from typing import Dict, Any, Tuple
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

logger = logging.getLogger(__name__)

class ResumeExporter:
    """Handles multi-format export for resumes."""

    def __init__(self):
        pass

    def export(self, resume: dict, export_format: str) -> Tuple[bytes, str, str]:
        """Export resume to requested format. Returns (bytes, media_type, filename)."""
        title = resume.get('title', 'resume')
        sanitized_title = "".join(c for c in title if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
        if not sanitized_title:
            sanitized_title = "Resume"

        if export_format == "pdf_formatted":
            content = self._generate_styled_pdf(resume)
            return content, "application/pdf", f"{sanitized_title}_formatted.pdf"

        elif export_format == "pdf_ats":
            content = self._generate_ats_pdf(resume)
            return content, "application/pdf", f"{sanitized_title}_ats.pdf"

        elif export_format == "docx":
            content = self._generate_docx(resume)
            return content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{sanitized_title}.docx"

        elif export_format == "txt":
            content = self._generate_txt(resume)
            return content, "text/plain", f"{sanitized_title}.txt"

        else:
            raise ValueError(f"Unsupported export format: {export_format}. Use pdf_formatted, pdf_ats, docx, or txt.")

    def _generate_txt(self, resume: dict) -> bytes:
        """Produce clean plaintext resume format."""
        pi = resume.get('personalInfo', {})
        lines = []

        lines.append("=" * 60)
        lines.append(f"{pi.get('fullName', 'RESUME').upper()}")
        contact = []
        if pi.get('email'): contact.append(pi['email'])
        if pi.get('phone'): contact.append(pi['phone'])
        if pi.get('location'): contact.append(pi['location'])
        if contact:
            lines.append(" | ".join(contact))
        links = []
        if pi.get('linkedin'): links.append(f"LinkedIn: {pi['linkedin']}")
        if pi.get('portfolio'): links.append(f"Portfolio: {pi['portfolio']}")
        if links:
            lines.append(" | ".join(links))
        lines.append("=" * 60)
        lines.append("")

        if resume.get('summary'):
            lines.append("PROFESSIONAL SUMMARY")
            lines.append("-" * 30)
            lines.append(resume['summary'])
            lines.append("")

        if resume.get('experience'):
            lines.append("WORK EXPERIENCE")
            lines.append("-" * 30)
            for exp in resume['experience']:
                lines.append(f"{exp.get('role', '')} - {exp.get('company', '')} ({exp.get('duration', '')})")
                if exp.get('description'):
                    lines.append(f"  {exp['description']}")
                for b in exp.get('bullets', []):
                    lines.append(f"  * {b}")
                if exp.get('skills'):
                    lines.append(f"  Technologies: {', '.join(exp['skills'])}")
                lines.append("")

        if resume.get('skills'):
            lines.append("TECHNICAL SKILLS")
            lines.append("-" * 30)
            skill_names = [s.get('skillName', '') if isinstance(s, dict) else str(s) for s in resume['skills']]
            lines.append(", ".join(filter(None, skill_names)))
            lines.append("")

        if resume.get('projects'):
            lines.append("KEY PROJECTS")
            lines.append("-" * 30)
            for proj in resume['projects']:
                lines.append(f"{proj.get('title', '')} ({proj.get('duration', '')})")
                if proj.get('description'):
                    lines.append(f"  {proj['description']}")
                for a in proj.get('achievements', []):
                    lines.append(f"  * {a}")
                if proj.get('technologies'):
                    lines.append(f"  Tech: {', '.join(proj['technologies'])}")
                lines.append("")

        if resume.get('education'):
            lines.append("EDUCATION")
            lines.append("-" * 30)
            for edu in resume['education']:
                gpa_str = f" - GPA: {edu.get('gpa')}" if edu.get('gpa') else ""
                lines.append(f"{edu.get('degree', '')} in {edu.get('field', '')} | {edu.get('institution', '')} ({edu.get('graduationDate', '')}){gpa_str}")
                if edu.get('relevantCoursework'):
                    lines.append(f"  Coursework: {', '.join(edu['relevantCoursework'])}")
                lines.append("")

        if resume.get('certifications'):
            lines.append("CERTIFICATIONS")
            lines.append("-" * 30)
            for cert in resume['certifications']:
                lines.append(f"* {cert.get('title', '')} - {cert.get('issuer', '')} ({cert.get('issueDate', '')})")
            lines.append("")

        return "\n".join(lines).encode('utf-8')

    def _generate_docx(self, resume: dict) -> bytes:
        """Create a Microsoft Word (.docx) resume with clean headings and bullets."""
        doc = Document()

        # Page margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.7)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.7)
            section.right_margin = Inches(0.7)

        pi = resume.get('personalInfo', {})
        # Name
        title_p = doc.add_paragraph()
        run = title_p.add_run(pi.get('fullName', 'Full Name'))
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(30, 41, 59)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Contact line
        contact_parts = []
        if pi.get('email'): contact_parts.append(pi['email'])
        if pi.get('phone'): contact_parts.append(pi['phone'])
        if pi.get('location'): contact_parts.append(pi['location'])
        if pi.get('linkedin'): contact_parts.append(pi['linkedin'])
        if pi.get('portfolio'): contact_parts.append(pi['portfolio'])

        if contact_parts:
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            crun = cp.add_run("  •  ".join(contact_parts))
            crun.font.size = Pt(9.5)
            crun.font.color.rgb = RGBColor(100, 116, 139)

        # Helper for headings
        def add_heading(text: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text.upper())
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(37, 99, 235)

        # Summary
        if resume.get('summary'):
            add_heading("Professional Summary")
            sp = doc.add_paragraph(resume['summary'])
            sp.paragraph_format.space_after = Pt(6)

        # Experience
        if resume.get('experience'):
            add_heading("Work Experience")
            for exp in resume['experience']:
                ep = doc.add_paragraph()
                ep.paragraph_format.space_before = Pt(4)
                ep.paragraph_format.space_after = Pt(2)
                r_role = ep.add_run(exp.get('role', 'Role'))
                r_role.bold = True
                ep.add_run(f" | {exp.get('company', 'Company')} ({exp.get('duration', '')})")

                if exp.get('description'):
                    dp = doc.add_paragraph(exp['description'])
                    dp.paragraph_format.space_after = Pt(2)

                for b in exp.get('bullets', []):
                    bp = doc.add_paragraph(b, style='List Bullet')
                    bp.paragraph_format.space_after = Pt(2)

        # Skills
        if resume.get('skills'):
            add_heading("Technical Competencies")
            skill_names = [s.get('skillName', '') if isinstance(s, dict) else str(s) for s in resume['skills']]
            sk_p = doc.add_paragraph(", ".join(filter(None, skill_names)))
            sk_p.paragraph_format.space_after = Pt(6)

        # Projects
        if resume.get('projects'):
            add_heading("Key Projects")
            for proj in resume['projects']:
                pp = doc.add_paragraph()
                pp.paragraph_format.space_before = Pt(4)
                pp.paragraph_format.space_after = Pt(2)
                r_title = pp.add_run(proj.get('title', 'Project'))
                r_title.bold = True
                pp.add_run(f" ({proj.get('duration', '')})")

                if proj.get('description'):
                    doc.add_paragraph(proj['description'])

                for a in proj.get('achievements', []):
                    doc.add_paragraph(a, style='List Bullet')

        # Education
        if resume.get('education'):
            add_heading("Education")
            for edu in resume['education']:
                edp = doc.add_paragraph()
                edp.paragraph_format.space_before = Pt(4)
                edp.paragraph_format.space_after = Pt(2)
                r_deg = edp.add_run(f"{edu.get('degree', '')} in {edu.get('field', '')}")
                r_deg.bold = True
                edp.add_run(f" — {edu.get('institution', '')} ({edu.get('graduationDate', '')})")

        # Certifications
        if resume.get('certifications'):
            add_heading("Certifications")
            for cert in resume['certifications']:
                cp = doc.add_paragraph(f"{cert.get('title', '')} — {cert.get('issuer', '')} ({cert.get('issueDate', '')})", style='List Bullet')
                cp.paragraph_format.space_after = Pt(2)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.read()

    def _generate_styled_pdf(self, resume: dict) -> bytes:
        """Generate visually appealing styled PDF using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'NameHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#0f172a'),
            alignment=1
        )
        contact_style = ParagraphStyle(
            'ContactSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#475569'),
            alignment=1
        )
        section_style = ParagraphStyle(
            'SectionHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#2563eb'),
            spaceBefore=10,
            spaceAfter=3
        )
        body_style = ParagraphStyle(
            'BodyTextCustom',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#1e293b')
        )
        bullet_style = ParagraphStyle(
            'BulletCustom',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            leftIndent=12,
            firstLineIndent=-8,
            textColor=colors.HexColor('#334155')
        )

        story = []
        pi = resume.get('personalInfo', {})

        # Name
        story.append(Paragraph(pi.get('fullName', 'Alex Morgan'), title_style))
        story.append(Spacer(1, 4))

        # Contacts
        contacts = []
        if pi.get('email'): contacts.append(pi['email'])
        if pi.get('phone'): contacts.append(pi['phone'])
        if pi.get('location'): contacts.append(pi['location'])
        if pi.get('linkedin'): contacts.append(pi['linkedin'])
        story.append(Paragraph(" &nbsp;|&nbsp; ".join(contacts), contact_style))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceBefore=2, spaceAfter=8))

        # Summary
        if resume.get('summary'):
            story.append(Paragraph("PROFESSIONAL SUMMARY", section_style))
            story.append(Paragraph(resume['summary'], body_style))
            story.append(Spacer(1, 6))

        # Experience
        if resume.get('experience'):
            story.append(Paragraph("EXPERIENCE", section_style))
            for exp in resume['experience']:
                header_line = f"<b>{exp.get('role', '')}</b> &nbsp;|&nbsp; {exp.get('company', '')} &nbsp;<i>({exp.get('duration', '')})</i>"
                story.append(Paragraph(header_line, body_style))
                for b in exp.get('bullets', []):
                    story.append(Paragraph(f"&bull; {b}", bullet_style))
                story.append(Spacer(1, 4))

        # Skills
        if resume.get('skills'):
            story.append(Paragraph("TECHNICAL SKILLS", section_style))
            skill_names = [s.get('skillName', '') if isinstance(s, dict) else str(s) for s in resume['skills']]
            story.append(Paragraph(", ".join(filter(None, skill_names)), body_style))
            story.append(Spacer(1, 6))

        # Projects
        if resume.get('projects'):
            story.append(Paragraph("PROJECTS", section_style))
            for proj in resume['projects']:
                p_head = f"<b>{proj.get('title', '')}</b> &nbsp;<i>({proj.get('duration', '')})</i>"
                story.append(Paragraph(p_head, body_style))
                for ach in proj.get('achievements', []):
                    story.append(Paragraph(f"&bull; {ach}", bullet_style))
                story.append(Spacer(1, 4))

        # Education
        if resume.get('education'):
            story.append(Paragraph("EDUCATION", section_style))
            for edu in resume['education']:
                e_line = f"<b>{edu.get('degree', '')} in {edu.get('field', '')}</b> &nbsp;|&nbsp; {edu.get('institution', '')} ({edu.get('graduationDate', '')})"
                story.append(Paragraph(e_line, body_style))
                story.append(Spacer(1, 4))

        # Certifications
        if resume.get('certifications'):
            story.append(Paragraph("CERTIFICATIONS", section_style))
            for c in resume['certifications']:
                c_line = f"&bull; {c.get('title', '')} &nbsp;&mdash;&nbsp; {c.get('issuer', '')} ({c.get('issueDate', '')})"
                story.append(Paragraph(c_line, bullet_style))
                story.append(Spacer(1, 2))

        doc.build(story)
        buffer.seek(0)
        return buffer.read()

    def _generate_ats_pdf(self, resume: dict) -> bytes:
        """Generate strictly single-column, standard ATS compliant PDF."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()
        header_style = ParagraphStyle(
            'ATSName',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.black,
            alignment=0
        )
        contact_style = ParagraphStyle(
            'ATSContact',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.black,
            alignment=0
        )
        section_style = ParagraphStyle(
            'ATSSection',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=colors.black,
            spaceBefore=10,
            spaceAfter=4
        )
        text_style = ParagraphStyle(
            'ATSText',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=colors.black
        )

        story = []
        pi = resume.get('personalInfo', {})

        story.append(Paragraph(pi.get('fullName', '').upper(), header_style))
        story.append(Spacer(1, 4))
        c_line = f"{pi.get('email', '')} | {pi.get('phone', '')} | {pi.get('location', '')}"
        story.append(Paragraph(c_line, contact_style))
        story.append(Spacer(1, 10))

        if resume.get('summary'):
            story.append(Paragraph("SUMMARY", section_style))
            story.append(Paragraph(resume['summary'], text_style))

        if resume.get('experience'):
            story.append(Paragraph("EXPERIENCE", section_style))
            for exp in resume['experience']:
                story.append(Paragraph(f"<b>{exp.get('company', '')}</b> - {exp.get('role', '')} ({exp.get('duration', '')})", text_style))
                for b in exp.get('bullets', []):
                    story.append(Paragraph(f"- {b}", text_style))
                story.append(Spacer(1, 4))

        if resume.get('skills'):
            story.append(Paragraph("SKILLS", section_style))
            names = [s.get('skillName', '') if isinstance(s, dict) else str(s) for s in resume['skills']]
            story.append(Paragraph(", ".join(filter(None, names)), text_style))

        if resume.get('projects'):
            story.append(Paragraph("PROJECTS", section_style))
            for p in resume['projects']:
                story.append(Paragraph(f"<b>{p.get('title', '')}</b> ({p.get('duration', '')})", text_style))
                for a in p.get('achievements', []):
                    story.append(Paragraph(f"- {a}", text_style))
                story.append(Spacer(1, 4))

        if resume.get('education'):
            story.append(Paragraph("EDUCATION", section_style))
            for e in resume['education']:
                story.append(Paragraph(f"{e.get('degree', '')} in {e.get('field', '')} - {e.get('institution', '')} ({e.get('graduationDate', '')})", text_style))

        if resume.get('certifications'):
            story.append(Paragraph("CERTIFICATIONS", section_style))
            for c in resume['certifications']:
                story.append(Paragraph(f"- {c.get('title', '')} ({c.get('issuer', '')}, {c.get('issueDate', '')})", text_style))

        doc.build(story)
        buffer.seek(0)
        return buffer.read()
