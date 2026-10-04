"""Unit tests for ATSAnalyzer service (Station 09: ATS Analyzer Card)."""

import os
import unittest
import asyncio
import tempfile
from docx import Document
from reportlab.pdfgen import canvas
from backend.services.ats_analyzer import ATSAnalyzer

class TestATSAnalyzer(unittest.TestCase):
    def setUp(self):
        self.loop = asyncio.new_event_loop()
        asyncio.set_event_loop(self.loop)
        self.analyzer = ATSAnalyzer()
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        self.loop.close()
        for root, dirs, files in os.walk(self.temp_dir, topdown=False):
            for name in files:
                try: os.remove(os.path.join(root, name))
                except: pass
            for name in dirs:
                try: os.rmdir(os.path.join(root, name))
                except: pass
        try: os.rmdir(self.temp_dir)
        except: pass

    def _create_sample_txt(self, content: str) -> str:
        path = os.path.join(self.temp_dir, "sample.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def _create_sample_docx(self, paragraphs, add_table=False) -> str:
        path = os.path.join(self.temp_dir, "sample.docx")
        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        if add_table:
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "Table Skill"
            t.cell(0, 1).text = "Docker"
            t.cell(1, 0).text = "Table Other"
            t.cell(1, 1).text = "AWS"
        doc.save(path)
        return path

    def _create_sample_pdf(self, lines) -> str:
        path = os.path.join(self.temp_dir, "sample.pdf")
        c = canvas.Canvas(path)
        y = 750
        for line in lines:
            c.drawString(50, y, line)
            y -= 20
        c.save()
        return path

    def test_parse_txt_file(self):
        async def _run():
            path = self._create_sample_txt("Python developer with AWS and Docker experience.")
            text = await self.analyzer._parse_resume(path)
            self.assertIn("Python developer", text)
            self.assertIn("AWS", text)
        self.loop.run_until_complete(_run())

    def test_parse_docx_file(self):
        async def _run():
            path = self._create_sample_docx(["Senior Engineer at Tech Corp", "Specialized in React, Python, and SQL"])
            text = await self.analyzer._parse_resume(path)
            self.assertIn("Senior Engineer", text)
            self.assertIn("React", text)
        self.loop.run_until_complete(_run())

    def test_parse_docx_with_table(self):
        async def _run():
            path = self._create_sample_docx(["Resume content"], add_table=True)
            text = await self.analyzer._parse_resume(path)
            self.assertIn("Docker", text)
            self.assertIn("AWS", text)
        self.loop.run_until_complete(_run())

    def test_parse_pdf_file(self):
        async def _run():
            path = self._create_sample_pdf(["Alex Doe Resume", "Machine Learning and Python Specialist", "Built FastAPI APIs"])
            text = await self.analyzer._parse_resume(path)
            self.assertIn("Machine Learning", text)
            self.assertIn("FastAPI", text)
        self.loop.run_until_complete(_run())

    def test_parse_unsupported_file_raises(self):
        async def _run():
            bad_path = os.path.join(self.temp_dir, "test.xyz")
            with open(bad_path, "w") as f:
                f.write("content")
            with self.assertRaises(ValueError):
                await self.analyzer._parse_resume(bad_path)
        self.loop.run_until_complete(_run())

    def test_formatting_check_clean_docx(self):
        async def _run():
            path = self._create_sample_docx(["Simple standard resume"])
            issues = await self.analyzer._check_formatting(path)
            self.assertEqual(len(issues), 0)
        self.loop.run_until_complete(_run())

    def test_formatting_check_docx_tables(self):
        async def _run():
            path = os.path.join(self.temp_dir, "tables.docx")
            doc = Document()
            doc.add_table(rows=2, cols=2)
            doc.add_table(rows=2, cols=2)
            doc.save(path)
            issues = await self.analyzer._check_formatting(path)
            self.assertTrue(any(i['type'] == 'tables_detected' for i in issues))
        self.loop.run_until_complete(_run())

    def test_analyze_keywords_found_and_missing(self):
        text = "Experienced in Python, React, AWS, and Docker. Implemented REST APIs and SQL databases."
        analysis = self.analyzer._analyze_keywords(text)
        found_names = [k['keyword'] for k in analysis['found']]
        self.assertIn('Python', found_names)
        self.assertIn('React', found_names)
        self.assertIn('AWS', found_names)
        self.assertIn('Docker', found_names)
        self.assertTrue(analysis['foundCount'] >= 4)
        self.assertTrue(len(analysis['missing']) > 0)
        self.assertTrue(analysis['matchPercentage'] > 0)

    def test_analyze_keywords_case_insensitivity(self):
        text = "Skilled in python, react, aws, and docker."
        analysis = self.analyzer._analyze_keywords(text)
        found_names = [k['keyword'] for k in analysis['found']]
        self.assertIn('Python', found_names)
        self.assertIn('React', found_names)

    def test_high_priority_missing_keywords(self):
        text = "Only know HTML and CSS."
        analysis = self.analyzer._analyze_keywords(text)
        high_missing = [k['keyword'] for k in analysis['missing'] if k['priority'] == 'high']
        self.assertIn('Python', high_missing)
        self.assertIn('Machine Learning', high_missing)

    def test_check_for_metrics_percentage(self):
        text = "Improved system performance by 45% and reduced latency by 30%."
        res = self.analyzer._check_for_metrics(text)
        self.assertTrue(res['metricsFound'] >= 2)
        self.assertIn('45%', res['examples'])
        self.assertIn('30%', res['examples'])

    def test_check_for_metrics_dollar(self):
        text = "Managed a budget of $500,000 and saved $50K annually."
        res = self.analyzer._check_for_metrics(text)
        self.assertTrue(res['metricsFound'] >= 2)

    def test_check_for_metrics_multipliers(self):
        text = "Scaled database throughput by 5x."
        res = self.analyzer._check_for_metrics(text)
        self.assertTrue(res['metricsFound'] >= 1)

    def test_simulate_ats_parsing_strips_bullets(self):
        raw = "• Led backend team \n * Developed microservices \n - Deployed to AWS"
        parsed = self.analyzer._simulate_ats_parsing(raw)
        self.assertNotIn("•", parsed)
        self.assertNotIn("*", parsed)
        self.assertIn("Led backend team", parsed)
        self.assertIn("Developed microservices", parsed)

    def test_compare_with_jd_high_match(self):
        resume = "Python, AWS, Docker, Kubernetes, React, FastAPI developer."
        jd = "Looking for a Python developer with AWS, Docker, and Kubernetes expertise."
        comp = self.analyzer._compare_with_jd(resume, jd)
        self.assertIsNotNone(comp)
        self.assertTrue(comp['matchPercentage'] >= 80)
        self.assertIn('Python', comp['matchingSkills'])
        self.assertIn('AWS', comp['matchingSkills'])

    def test_compare_with_jd_low_match(self):
        resume = "Graphic designer skilled in Photoshop and Figma."
        jd = "Requirements: Python, Machine Learning, TensorFlow, PyTorch, Docker, Kubernetes."
        comp = self.analyzer._compare_with_jd(resume, jd)
        self.assertIsNotNone(comp)
        self.assertTrue(comp['matchPercentage'] < 30)
        self.assertTrue(len(comp['missingSkills']) >= 4)

    def test_calculate_ats_score_no_jd(self):
        keywords = {'matchPercentage': 20.0}
        metrics = {'metricsFound': 4, 'hasMetrics': True}
        formatting = []
        score = self.analyzer._calculate_ats_score(formatting, keywords, metrics, None)
        self.assertIn('overall', score)
        self.assertIn('formatting', score)
        self.assertIn('keywords', score)
        self.assertIn('metrics', score)
        self.assertIn('readability', score)
        self.assertIn('parsing', score)
        self.assertTrue(score['overall'] >= 60)

    def test_calculate_ats_score_with_jd(self):
        keywords = {'matchPercentage': 25.0}
        metrics = {'metricsFound': 5, 'hasMetrics': True}
        formatting = []
        jd_comp = {'matchPercentage': 90.0}
        score = self.analyzer._calculate_ats_score(formatting, keywords, metrics, jd_comp)
        self.assertTrue(score['overall'] >= 75)

    def test_generate_recommendations_categories(self):
        formatting = [{'severity': 'error', 'suggestion': 'Fix unreadable formatting'}]
        keywords = {'missing': [{'keyword': 'Python', 'priority': 'high'}], 'hasMetrics': False}
        jd_comp = {'missingSkills': ['Kubernetes']}
        score = {'overall': 65, 'metrics': 50}
        recs = self.analyzer._generate_recommendations(formatting, keywords, jd_comp, score)
        self.assertIn('immediate', recs)
        self.assertIn('shortTerm', recs)
        self.assertIn('longTerm', recs)
        self.assertTrue(len(recs['immediate']) >= 2)
        self.assertTrue(len(recs['shortTerm']) >= 2)
        self.assertTrue(len(recs['longTerm']) >= 1)

    def test_full_analyze_txt_pipeline(self):
        async def _run():
            content = """
            ALEX MORGAN
            alex@example.com | 555-1234
            SUMMARY: Software Engineer experienced with Python, React, FastAPI, AWS, and Docker.
            EXPERIENCE:
            Senior Engineer at FlowTech:
            - Architected REST APIs, improving throughput by 40% and saving $35K.
            - Managed Kubernetes clusters and PostgreSQL databases.
            EDUCATION:
            BS Computer Science
            SKILLS: Python, React, Docker, AWS, SQL, FastAPI, Git, CI/CD
            """
            path = self._create_sample_txt(content)
            res = await self.analyzer.analyze(path, job_description="Looking for Python and AWS engineer")
            self.assertTrue(res['success'])
            self.assertIn('analysisId', res)
            self.assertTrue(res['atsScore']['overall'] >= 70)
            self.assertTrue(res['jdComparison']['matchPercentage'] >= 80)
        self.loop.run_until_complete(_run())

    def test_analyze_resume_object(self):
        async def _run():
            resume_obj = {
                'personalInfo': {'fullName': 'Taylor Swift', 'email': 'taylor@example.com'},
                'summary': 'Full stack developer proficient in Python and React.',
                'experience': [{'company': 'Stripe', 'role': 'Developer', 'bullets': ['Increased conversions by 25%']}],
                'skills': [{'skillName': 'Python'}, {'skillName': 'React'}, {'skillName': 'Docker'}]
            }
            res = await self.analyzer.analyze_resume_object(resume_obj)
            self.assertTrue(res['success'])
            self.assertTrue(res['atsScore']['overall'] >= 50)
            self.assertIn('immediate', res['recommendations'])
        self.loop.run_until_complete(_run())

    def test_extract_skills_from_text(self):
        text = "Deep understanding of TypeScript, GraphQL, Next.js, and Redis."
        skills = self.analyzer._extract_skills_from_text(text)
        self.assertIn("TypeScript", skills)
        self.assertIn("GraphQL", skills)
        self.assertIn("Next.js", skills)
        self.assertIn("Redis", skills)

    def test_formatting_check_empty_pdf(self):
        async def _run():
            path = os.path.join(self.temp_dir, "empty.pdf")
            c = canvas.Canvas(path)
            c.save()
            issues = await self.analyzer._check_formatting(path)
            self.assertIsInstance(issues, list)
        self.loop.run_until_complete(_run())

    def test_metrics_no_numbers(self):
        text = "Responsible for daily operations and attending meetings."
        res = self.analyzer._check_for_metrics(text)
        self.assertEqual(res['metricsFound'], 0)
        self.assertFalse(res['hasMetrics'])

    def test_resume_dict_to_text_serialization(self):
        resume_dict = {
            'personalInfo': {'fullName': 'Jordan Lee', 'email': 'jordan@email.com'},
            'summary': 'Machine learning researcher.',
            'experience': [{'company': 'DeepMind', 'role': 'Scientist', 'bullets': ['Invented new transformer model']}],
            'education': [{'degree': 'PhD', 'institution': 'MIT'}],
            'skills': [{'skillName': 'PyTorch'}, {'skillName': 'Python'}],
            'projects': [{'title': 'AlphaFold Open', 'technologies': ['Python', 'PyTorch']}],
            'certifications': [{'title': 'Deep Learning Specialization'}]
        }
        text = self.analyzer._resume_dict_to_text(resume_dict)
        self.assertIn("Jordan Lee", text)
        self.assertIn("Machine learning researcher", text)
        self.assertIn("DeepMind", text)
        self.assertIn("PyTorch", text)
        self.assertIn("AlphaFold Open", text)

    def test_warnings_generated_on_short_resume(self):
        async def _run():
            path = self._create_sample_txt("Too short.")
            res = await self.analyzer.analyze(path)
            self.assertTrue(len(res['warnings']) >= 1)
        self.loop.run_until_complete(_run())

    def test_score_bounds(self):
        for score_val in [0.0, 50.0, 100.0]:
            kw = {'matchPercentage': score_val}
            m = {'metricsFound': 3, 'hasMetrics': True}
            scores = self.analyzer._calculate_ats_score([], kw, m, None)
            self.assertTrue(10.0 <= scores['overall'] <= 100.0)

if __name__ == '__main__':
    unittest.main()
